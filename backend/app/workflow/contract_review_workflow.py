"""Contract review workflow orchestration."""

import asyncio
import hashlib
import json
from datetime import datetime
from io import BytesIO

from langgraph.graph import END, StateGraph
from loguru import logger
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.common.snowflake import generate_id
from app.core.database import async_session_factory
from app.models.contract_document import ContractDocument
from app.models.workflow_task import WorkflowTask
from app.repositories.contract_document_repo import ContractDocumentRepository
from app.repositories.workflow_task_repo import WorkflowTaskRepository
from app.services.approval.approval_service import ApprovalService
from app.services.risk.risk_service import RiskReviewService
from app.workflow.checkpointer import mysql_checkpointer
from app.workflow.workflow_nodes import (
    block_task_node,
    complete_task_node,
    download_contract_node,
    extract_fields_node,
    llm_review_node,
    parse_document_node,
    persist_task_state,
    pull_approval_node,
    review_risks_node,
    write_comment_node,
)
from app.workflow.workflow_state import WorkflowState


class RetryConfig:
    MAX_RETRIES = 3
    BASE_DELAY_SECONDS = 5
    MAX_DELAY_SECONDS = 300
    BACKOFF_MULTIPLIER = 2

    RETRYABLE_ERRORS = {
        "ConnectionError",
        "TimeoutError",
        "HTTPError",
        "RequestException",
        "TemporaryError",
    }

    BLOCKABLE_ERRORS = {
        "NoAttachmentError",
        "MissingDocIdError",
        "MissingParseTextError",
        "ValueError",
    }

    @classmethod
    def get_delay(cls, retry_count: int) -> int:
        return min(cls.BASE_DELAY_SECONDS * (cls.BACKOFF_MULTIPLIER**retry_count), cls.MAX_DELAY_SECONDS)

    @classmethod
    def is_retryable(cls, error_type: str) -> bool:
        return error_type in cls.RETRYABLE_ERRORS

    @classmethod
    def should_block(cls, error_type: str) -> bool:
        return error_type in cls.BLOCKABLE_ERRORS


def _route_after_pull(state: WorkflowState) -> str:
    return "download_contract" if state.get("approval_exists") else "block_task"


def _route_after_download(state: WorkflowState) -> str:
    if state.get("doc_id"):
        return "parse_document"
    return _handle_error_route(state, "download_contract")


def _route_after_parse(state: WorkflowState) -> str:
    if state.get("parse_status") == "parsed":
        return "extract_fields"
    return _handle_error_route(state, "parse_document")


def _route_after_review(state: WorkflowState) -> str:
    return "write_comment" if state.get("has_risks") else "complete_task"


def _handle_error_route(state: WorkflowState, current_node: str) -> str:
    error_type = state.get("error_type", "")
    retry_count = state.get("retry_count", 0)
    max_retries = state.get("max_retries", RetryConfig.MAX_RETRIES)

    if RetryConfig.should_block(error_type):
        return "block_task"
    if RetryConfig.is_retryable(error_type) and retry_count < max_retries:
        logger.info(f"Retry node {current_node} | retry={retry_count + 1}/{max_retries}")
        return current_node
    return "block_task"


def build_contract_review_graph() -> StateGraph:
    workflow = StateGraph(WorkflowState)
    workflow.add_node("pull_approval", _make_node(pull_approval_node))
    workflow.add_node("download_contract", _make_node(download_contract_node))
    workflow.add_node("parse_document", _make_node(parse_document_node))
    workflow.add_node("extract_fields", _make_node(extract_fields_node))
    workflow.add_node("review_risks", _make_node(review_risks_node))
    workflow.add_node("write_comment", _make_node(write_comment_node))
    workflow.add_node("complete_task", _make_node(complete_task_node))
    workflow.add_node("block_task", _make_node(block_task_node))

    workflow.set_entry_point("pull_approval")
    workflow.add_conditional_edges(
        "pull_approval",
        _route_after_pull,
        {"download_contract": "download_contract", "block_task": "block_task"},
    )
    workflow.add_conditional_edges(
        "download_contract",
        _route_after_download,
        {
            "parse_document": "parse_document",
            "download_contract": "download_contract",
            "block_task": "block_task",
        },
    )
    workflow.add_conditional_edges(
        "parse_document",
        _route_after_parse,
        {
            "extract_fields": "extract_fields",
            "parse_document": "parse_document",
            "block_task": "block_task",
        },
    )
    workflow.add_edge("extract_fields", "review_risks")
    workflow.add_conditional_edges(
        "review_risks",
        _route_after_review,
        {"write_comment": "write_comment", "complete_task": "complete_task"},
    )
    workflow.add_edge("write_comment", "complete_task")
    workflow.add_edge("complete_task", END)
    workflow.add_edge("block_task", END)
    return workflow


def _make_node(node_func):
    async def wrapper(state: WorkflowState) -> dict:
        retry_count = state.get("retry_count", 0)
        if retry_count > 0:
            await asyncio.sleep(RetryConfig.get_delay(retry_count))

        async with async_session_factory() as db:
            result = await node_func(state, db)
            if result.get("error_message") and RetryConfig.is_retryable(result.get("error_type", "")):
                result["retry_count"] = retry_count + 1

            merged = {**state, **result}
            await persist_task_state(merged, db)
            await db.commit()
        return result

    return wrapper


class ContractReviewWorkflow:
    def __init__(self):
        self._graph = build_contract_review_graph()
        self._compiled = None
        self._running_tasks: dict[str, asyncio.Task] = {}  # 持有异步任务引用，防止被 GC 回收

    @property
    def compiled(self):
        if self._compiled is None:
            self._compiled = self._graph.compile(checkpointer=mysql_checkpointer)
            logger.info("Contract review workflow graph compiled")
        return self._compiled

    async def create_task(self, db: AsyncSession, approval_order_id: int) -> str:
        task_id = f"TASK_{generate_id()}"
        task_repo = WorkflowTaskRepository(db)
        task = WorkflowTask(
            id=generate_id(),
            task_id=task_id,
            approval_order_id=approval_order_id,
            workflow_name="contract_review",
            status="pending",
            current_node="",
            retry_count=0,
            started_at=datetime.now(),
        )
        await task_repo.create(task)
        logger.info(f"Review task created | task_id={task_id} | approval_order_id={approval_order_id}")
        return task_id

    async def run(self, task_id: str, approval_order_id: int) -> dict:
        config = {"configurable": {"thread_id": task_id}}
        initial_state: WorkflowState = {
            "task_id": task_id,
            "approval_order_id": approval_order_id,
            "workflow_name": "contract_review",
            "approval_exists": False,
            "doc_exists": False,
            "parse_status": "pending",
            "page_count": 0,
            "is_scanned": False,
            "field_extracted": False,
            "extract_confidence": 0.0,
            "risk_results": [],
            "risk_score": {},
            "risk_count": 0,
            "has_risks": False,
            "comment_written": False,
            "current_node": "",
            "error_message": None,
            "error_type": None,
            "retry_count": 0,
            "max_retries": RetryConfig.MAX_RETRIES,
            "is_blocked": False,
            "block_reason": None,
            "messages": [],
            "started_at": datetime.now().isoformat(),
            "completed_at": None,
            "duration_ms": None,
            "node_statuses": {},
            "failed_node": None,
        }

        final_state = None
        async for event in self.compiled.astream(initial_state, config, stream_mode="values"):
            final_state = event

        async with async_session_factory() as db:
            final_state = await self._finalize_review_outputs(db, final_state or {})
            await persist_task_state(final_state, db)
            await db.commit()
        return self._format_result(final_state)

    async def resume(self, task_id: str) -> dict:
        config = {"configurable": {"thread_id": task_id}}
        try:
            current = self.compiled.get_state(config)
            resume_input = current.values if current and current.values else {}
        except Exception:
            resume_input = {}

        final_state = None
        async for event in self.compiled.astream(resume_input, config, stream_mode="values"):
            final_state = event

        if final_state:
            async with async_session_factory() as db:
                final_state = await self._finalize_review_outputs(db, final_state)
                await persist_task_state(final_state, db)
                await db.commit()

        return self._format_result(final_state or {})

    async def run_background(self, task_id: str, approval_order_id: int) -> None:
        """异步启动工作流任务，持有引用防止被 GC 回收"""
        # 先更新任务状态为 running，前端立即可见
        async with async_session_factory() as db:
            task_repo = WorkflowTaskRepository(db)
            task = await task_repo.get_by_task_id(task_id)
            if task:
                task.status = "running"
                await task_repo.update(task)
                await db.commit()

        # 创建异步任务并持有引用
        bg_task = asyncio.create_task(
            self._run_with_error_handling(task_id, approval_order_id)
        )
        self._running_tasks[task_id] = bg_task
        # 任务完成后自动清理引用
        bg_task.add_done_callback(lambda _: self._running_tasks.pop(task_id, None))

    async def _run_with_error_handling(self, task_id: str, approval_order_id: int) -> None:
        try:
            await self.run(task_id, approval_order_id)
        except Exception as exc:
            logger.exception(f"Workflow execution failed | task_id={task_id} | error={exc}")
            async with async_session_factory() as db:
                task_repo = WorkflowTaskRepository(db)
                task = await task_repo.get_by_task_id(task_id)
                if task:
                    task.status = "failed"
                    task.error_message = str(exc)
                    await task_repo.update(task)
                    await db.commit()

    @staticmethod
    def _format_result(state: dict) -> dict:
        return {
            "task_id": state.get("task_id"),
            "status": "blocked" if state.get("is_blocked") else "success",
            "current_node": state.get("current_node"),
            "has_risks": state.get("has_risks", False),
            "risk_count": state.get("risk_count", 0),
            "risk_level": state.get("risk_score", {}).get("overall_level", "NONE"),
            "comment_written": state.get("comment_written", False),
            "is_blocked": state.get("is_blocked", False),
            "block_reason": state.get("block_reason"),
            "error_message": state.get("error_message"),
            "duration_ms": state.get("duration_ms"),
            "node_statuses": state.get("node_statuses", {}),
            "failed_node": state.get("failed_node"),
            "approval_order_id": state.get("approval_order_id"),
            "review_report_doc_id": state.get("review_report_doc_id"),
        }

    async def run_direct_review(self, task_id: str, doc_id: str) -> dict:
        started_at = datetime.now().isoformat()

        async with async_session_factory() as db:
            doc_repo = ContractDocumentRepository(db)
            task_repo = WorkflowTaskRepository(db)
            approval_service = ApprovalService(db)

            doc = await doc_repo.get_by_doc_id(doc_id)
            if not doc:
                raise ValueError(f"Document not found: {doc_id}")

            approval_order_id = doc.approval_order_id
            if approval_order_id == 0:
                order = await approval_service.ensure_order_for_document(doc_id, doc.file_name)
                await db.flush()
                approval_order_id = order.id
                doc.approval_order_id = approval_order_id
                await doc_repo.update(doc)
                await db.flush()
                logger.info(
                    f"Approval order created for direct review | approval_order_id={approval_order_id} | doc_id={doc_id}"
                )

            task = WorkflowTask(
                id=generate_id(),
                task_id=task_id,
                approval_order_id=approval_order_id,
                workflow_name="direct_review",
                status="running",
                current_node="parse_document",
                retry_count=0,
                started_at=datetime.now(),
            )
            await task_repo.create(task)

            state: WorkflowState = {
                "task_id": task_id,
                "approval_order_id": approval_order_id,
                "workflow_name": "direct_review",
                "doc_id": doc_id,
                "approval_exists": True,
                "doc_exists": True,
                "parse_status": "parsed",
                "parse_text": doc.parse_text or "",
                "page_count": 0,
                "is_scanned": bool(doc.is_scanned),
                "field_extracted": False,
                "extract_confidence": 0.0,
                "risk_results": [],
                "risk_score": {},
                "risk_count": 0,
                "has_risks": False,
                "comment_written": False,
                "current_node": "parse_document",
                "error_message": None,
                "error_type": None,
                "retry_count": 0,
                "max_retries": RetryConfig.MAX_RETRIES,
                "is_blocked": False,
                "block_reason": None,
                "messages": [],
                "started_at": started_at,
                "completed_at": None,
                "duration_ms": None,
                "node_statuses": {},
                "failed_node": None,
            }

            try:
                state.update(await parse_document_node(state, db))
                state["current_node"] = "parse_document"
                await persist_task_state(state, db)
                await db.commit()
                if state.get("parse_status") != "parsed":
                    state.update(
                        is_blocked=True,
                        block_reason=state.get("error_message", "Document parse failed"),
                        current_node="blocked",
                    )
                    await persist_task_state(state, db)
                    await db.commit()
                    return self._format_result(state)

                state.update(await extract_fields_node(state, db))
                state["current_node"] = "extract_fields"
                await persist_task_state(state, db)
                await db.commit()

                state.update(await review_risks_node(state, db))
                state["current_node"] = "review_risks"
                await persist_task_state(state, db)
                await db.commit()

                state.update(await llm_review_node(state, db))
                state["current_node"] = "llm_review"
                await persist_task_state(state, db)
                await db.commit()

                if state.get("has_risks"):
                    state.update(await write_comment_node(state, db))
                    state["current_node"] = "write_comment"
                    await persist_task_state(state, db)
                    await db.commit()

                state.update(await complete_task_node(state, db))
                state = await self._finalize_review_outputs(db, state, source_doc=doc)
                await persist_task_state(state, db)
                await db.commit()
            except Exception:
                await db.rollback()
                raise

        result = self._format_result(state)
        logger.info(
            f"Direct review completed | task_id={task_id} | risks={result['risk_count']} | level={result['risk_level']}"
        )
        return result

    async def _finalize_review_outputs(
        self,
        db: AsyncSession,
        state: dict,
        source_doc: ContractDocument | None = None,
    ) -> dict:
        if not state or state.get("is_blocked"):
            return state

        approval_order_id = state.get("approval_order_id")
        if not approval_order_id:
            return state

        risk_results = state.get("risk_results", []) or []
        risk_service = RiskReviewService(db)
        await risk_service.replace_results_for_approval(approval_order_id, risk_results)

        if source_doc is None:
            source_doc = await self._resolve_source_document(db, state)

        if source_doc and risk_results:
            state["review_report_doc_id"] = await self._generate_review_report(
                db,
                source_doc,
                risk_results,
                state.get("risk_score", {}),
            )

        return state

    async def _resolve_source_document(self, db: AsyncSession, state: dict) -> ContractDocument | None:
        doc_repo = ContractDocumentRepository(db)
        doc_id = state.get("doc_id")
        if doc_id:
            doc = await doc_repo.get_by_doc_id(doc_id)
            if doc:
                return doc

        docs = await doc_repo.list_by_approval_order(state.get("approval_order_id", 0))
        return docs[0] if docs else None

    async def _generate_review_report(
        self,
        db: AsyncSession,
        doc: ContractDocument,
        risk_results: list[dict],
        risk_score: dict,
    ) -> str:
        from app.infrastructure.minio_client import minio_client

        risk_level = (risk_score.get("overall_level") or "LOW").upper()
        report_name = self._build_report_file_name(doc.file_name, risk_level)
        report_text = self._build_report_text(doc.file_name, risk_results, risk_score, risk_level)
        report_bytes = self._render_report_bytes(doc.file_type, report_text)
        file_md5 = hashlib.md5(report_bytes).hexdigest()

        report_doc = await self._find_existing_report(db, doc.approval_order_id, report_name)
        if report_doc is None:
            report_doc = ContractDocument(
                id=generate_id(),
                doc_id=f"DOC_{generate_id()}",
                approval_order_id=doc.approval_order_id,
                file_name=report_name,
                file_md5=file_md5,
                minio_path="",
                file_type=doc.file_type,
                file_size=len(report_bytes),
                is_scanned=False,
                parse_status="parsed",
                parse_text=report_text,
            )
            db.add(report_doc)
            await db.flush()
        else:
            report_doc.file_name = report_name
            report_doc.file_md5 = file_md5
            report_doc.file_type = doc.file_type
            report_doc.file_size = len(report_bytes)
            report_doc.parse_status = "parsed"
            report_doc.parse_text = report_text
            await db.flush()

        minio_path = f"contracts/{report_doc.doc_id}/{report_name}"
        await minio_client.upload(
            object_name=minio_path,
            data=report_bytes,
            content_type=self._resolve_content_type(doc.file_type, report_name),
        )
        report_doc.minio_path = minio_path
        await db.flush()
        return report_doc.doc_id

    @staticmethod
    def _build_report_file_name(original_file_name: str, risk_level: str) -> str:
        name, dot, ext = original_file_name.rpartition(".")
        if dot:
            return f"{risk_level}_{name}.{ext}"
        return f"{risk_level}_{original_file_name}"

    @staticmethod
    def _build_report_text(
        original_file_name: str,
        risk_results: list[dict],
        risk_score: dict,
        risk_level: str,
    ) -> str:
        lines = [
            "Contract Risk Review Report",
            f"Original file: {original_file_name}",
            f"Review time: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
            f"Overall level: {risk_level}",
            f"Risk score: {risk_score.get('score', 0)}",
            f"Risk count: {risk_score.get('total_count', len(risk_results))}",
            "",
        ]
        for index, item in enumerate(risk_results, start=1):
            lines.append(
                f"{index}. [{item.get('risk_level', 'LOW')}] "
                f"{item.get('rule_name') or item.get('risk_name') or 'Risk Item'}"
            )
            lines.append(f"Description: {item.get('risk_description', '')}")
            if item.get("suggestion"):
                lines.append(f"Suggestion: {item.get('suggestion')}")
            if item.get("source_text"):
                lines.append(f"Source text: {item.get('source_text')}")
            lines.append("")
        return "\n".join(lines)

    @staticmethod
    def _resolve_content_type(file_type: str, file_name: str) -> str:
        if file_type == "pdf":
            return "application/pdf"
        if file_type == "word":
            if file_name.lower().endswith(".doc"):
                return "application/msword"
            return "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        return "application/octet-stream"

    @staticmethod
    def _render_report_bytes(file_type: str, report_text: str) -> bytes:
        if file_type == "word":
            from docx import Document as DocxDocument

            document = DocxDocument()
            document.add_heading("Contract Risk Review Report", level=1)
            for line in report_text.splitlines():
                document.add_paragraph(line)
            buffer = BytesIO()
            document.save(buffer)
            return buffer.getvalue()

        return report_text.encode("utf-8")

    @staticmethod
    async def _find_existing_report(
        db: AsyncSession,
        approval_order_id: int,
        report_name: str,
    ) -> ContractDocument | None:
        result = await db.execute(
            select(ContractDocument).where(
                ContractDocument.approval_order_id == approval_order_id,
                ContractDocument.file_name == report_name,
            )
        )
        return result.scalar_one_or_none()


contract_review_workflow = ContractReviewWorkflow()
